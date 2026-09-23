import openpyxl
from datetime import datetime, date

SRC = '/mnt/user-data/uploads/戰情表資料庫.xlsx'
wb = openpyxl.load_workbook(SRC, data_only=True)

def sql_str(v):
    if v is None:
        return 'NULL'
    s = str(v)
    return "'" + s.replace("'", "''") + "'"

def sql_num(v):
    if v is None or v == '':
        return 'NULL'
    try:
        return str(float(v))
    except Exception:
        return 'NULL'

def sql_date(v):
    if v is None:
        return 'NULL'
    if isinstance(v, (datetime, date)):
        return sql_str(v.strftime('%Y-%m-%d'))
    return sql_str(str(v)[:10])

def sql_ts(v):
    if v is None:
        return 'NULL'
    if isinstance(v, (datetime, date)):
        return sql_str(v.isoformat())
    return sql_str(str(v))

out = []
out.append("-- =====================================================================")
out.append("-- 太綺戰情室 — 資料匯入 seed.sql（由 Excel 自動產生）")
out.append("-- 請先執行過 database/schema.sql 建好資料表，再執行這份")
out.append("-- =====================================================================")
out.append("BEGIN;")
out.append("")

# ---------------------------------------------------------------------
# 1) 案場主檔 — 新格式(P0001~P0008) 目前完全沒有主檔資料，先建立「待確認」佔位資料
#    避免後面的支出/發票/月度請款筆記錄因為外鍵找不到案場而失敗。
#    ⚠️ 請上線前務必到「案場列表」把名稱/業主/合約金額改成正確資料。
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 1) 案場主檔：P0001~P0008 佔位資料（原始檔案中的案場主檔分頁是空的）")
out.append("--    ⚠️ 請務必之後在畫面上把名稱/業主/合約金額改成正確內容")
out.append("-- ---------------------------------------------------------------")
placeholder_ids = ['P0001','P0002','P0003','P0004','P0005','P0006','P0007','P0008']
for pid in placeholder_ids:
    out.append(
        f"INSERT INTO \"案場主檔\"(\"案場ID\",\"案場名稱\",\"業主\",\"主約未稅金額\",\"主約稅額\",\"主約含稅金額\",\"預設付款條件\",\"預設付款方式\",\"狀態\",\"開案日期\",\"備註\") "
        f"VALUES ({sql_str(pid)}, {sql_str('⚠️ 請填寫案場名稱 (' + pid + ')')}, NULL, 0, 0, 0, NULL, NULL, '進行中', NULL, "
        f"{sql_str('系統自動建立的佔位資料，原始資料庫的案場主檔分頁沒有這張案場的紀錄，請手動補上正確名稱/業主/合約金額')}) "
        f"ON CONFLICT (\"案場ID\") DO NOTHING;"
    )
out.append("")

# ---------------------------------------------------------------------
# 2) 廠商主檔
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 2) 廠商主檔")
out.append("-- ---------------------------------------------------------------")
ws = wb['廠商主檔']
rows = list(ws.iter_rows(min_row=2, values_only=True))
def clean_phone(v):
    if v in (None, '-', ''):
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

for r in rows:
    vid, taxid, name, addr, contact, phone, created, note = r
    if not vid:
        continue
    out.append(
        f"INSERT INTO \"廠商主檔\"(\"廠商ID\",\"統一編號\",\"廠商名稱\",\"地址\",\"聯絡人\",\"電話\",\"建立日期\",\"備註\") VALUES "
        f"({sql_str(vid)}, {sql_str(str(int(taxid)) if isinstance(taxid,(int,float)) else taxid)}, {sql_str(name)}, "
        f"{sql_str(None if addr in ('-',None) else addr)}, {sql_str(None if contact in ('-',None) else contact)}, "
        f"{sql_str(clean_phone(phone))}, {sql_ts(created)}, {sql_str(note)}) "
        f"ON CONFLICT (\"廠商ID\") DO NOTHING;"
    )
out.append("")

# ---------------------------------------------------------------------
# 3) 月度請款（新格式既有的 1 筆）
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 3) 月度請款")
out.append("-- ---------------------------------------------------------------")
ws = wb['月度請款']
for r in ws.iter_rows(min_row=2, values_only=True):
    bid, pid, ym, est, method, cashpct, checkpct, note = r
    if not bid:
        continue
    ym_str = ym.strftime('%Y-%m') if isinstance(ym, (datetime, date)) else str(ym)[:7]
    out.append(
        f"INSERT INTO \"月度請款\"(\"請款ID\",\"案場ID\",\"年月\",\"預估請款金額\",\"收入方式\",\"現金比例\",\"支票比例\",\"備註\") VALUES "
        f"({sql_str(bid)}, {sql_str(pid)}, {sql_str(ym_str)}, {sql_num(est)}, {sql_str(method)}, {sql_num(cashpct)}, {sql_num(checkpct)}, {sql_str(note)}) "
        f"ON CONFLICT (\"請款ID\") DO NOTHING;"
    )
out.append("")

# ---------------------------------------------------------------------
# 4) 發票登記
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 4) 發票登記")
out.append("-- ---------------------------------------------------------------")
ws = wb['發票登記']
for r in ws.iter_rows(min_row=2, values_only=True):
    iid, pid, typ, taxid, idate, ino, item, amt, note = r
    if not iid:
        continue
    out.append(
        f"INSERT INTO \"發票登記\"(\"發票ID\",\"案場ID\",\"類型\",\"統一編號\",\"發票日期\",\"發票號碼\",\"項目\",\"金額(含稅)\",\"備註\") VALUES "
        f"({sql_str(iid)}, {sql_str(pid)}, {sql_str(typ)}, {sql_str(taxid)}, {sql_date(idate)}, {sql_str(ino)}, {sql_str(item)}, {sql_num(amt)}, {sql_str(note)}) "
        f"ON CONFLICT (\"發票ID\") DO NOTHING;"
    )
out.append("")

# ---------------------------------------------------------------------
# 5) 支出登記
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 5) 支出登記")
out.append("-- ---------------------------------------------------------------")
ws = wb['支出登記']
for r in ws.iter_rows(min_row=2, values_only=True):
    eid, pid, typ, edate, vendor, ino, item, amt, worker, task, qty, subtotal, tax_inc, tax_health, actual, note = r
    if not eid:
        continue
    out.append(
        f"INSERT INTO \"支出登記\"(\"支出ID\",\"案場ID\",\"類型\",\"日期\",\"廠商ID\",\"發票號碼\",\"項目\",\"金額(含稅)\",\"點工人員ID\",\"工項\",\"數量\",\"小計\",\"代扣所得稅\",\"代扣二代健保\",\"實際支付金額\",\"備註\") VALUES "
        f"({sql_str(eid)}, {sql_str(pid)}, {sql_str(typ)}, {sql_date(edate)}, {sql_str(vendor)}, {sql_str(ino)}, {sql_str(item)}, {sql_num(amt)}, "
        f"{sql_str(worker)}, {sql_str(task)}, {sql_num(qty)}, {sql_num(subtotal)}, {sql_num(tax_inc)}, {sql_num(tax_health)}, {sql_num(actual)}, {sql_str(note)}) "
        f"ON CONFLICT (\"支出ID\") DO NOTHING;"
    )
out.append("")

# ---------------------------------------------------------------------
# 6) 變更記錄（稽核紀錄，原封不動搬過去）
# ---------------------------------------------------------------------
out.append("-- ---------------------------------------------------------------")
out.append("-- 6) 變更記錄")
out.append("-- ---------------------------------------------------------------")
ws = wb['變更記錄']
for r in ws.iter_rows(min_row=2, values_only=True):
    t, who, action, target, tid, note = r
    if not t:
        continue
    out.append(
        f"INSERT INTO \"變更記錄\"(\"時間\",\"操作人\",\"動作\",\"目標分頁\",\"目標ID\",\"備註\") VALUES "
        f"({sql_ts(t)}, {sql_str(who)}, {sql_str(action)}, {sql_str(target)}, {sql_str(tid)}, {sql_str(note)});"
    )
out.append("")

out.append("COMMIT;")

with open('/home/claude/taichi/database/seed_current.sql', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))

print('done, lines:', len(out))
